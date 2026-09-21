from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "use-cases" / "customer-use-case-specifications.docx"


TRACEABILITY = [
    ("UC-01", "Đăng ký tài khoản", "FR-01 / #23", "III-00", "III-01-registration.puml"),
    ("UC-02", "Đăng nhập, làm mới phiên và đăng xuất", "FR-02 / #24", "III-00", "III-02-login-refresh-logout.puml"),
    ("UC-03", "Quản lý hồ sơ và sở thích", "FR-03 / #25", "III-00", "III-03-profile-preference.puml"),
    ("UC-04", "Gửi feedback và rating sản phẩm", "FR-04 / #26", "III-00", "III-04-feedback-rating-verified-purchase.puml"),
    ("UC-05", "Trả lời khảo sát", "FR-05 / #27", "III-00", "III-05-survey-response.puml"),
    ("UC-06", "Xem, tìm kiếm và lọc sản phẩm", "FR-06 / #28", "III-00", "III-06-product-catalog.puml"),
    ("UC-07", "Giỏ hàng, đặt hàng và thanh toán payOS", "FR-07 / #21, #29", "III-00", "III-07-shopping-cart-order-payos.puml"),
    ("UC-08", "Xem lịch sử và trạng thái đơn hàng", "FR-08 / #30", "III-00", "III-08-order-history-status-tracking.puml"),
]


USE_CASES = [
    {
        "id": "UC-01",
        "title": "Đăng ký tài khoản",
        "actors": "Customer",
        "support": "Customer Portal, Account API, CRM Database",
        "trigger": "Customer chọn chức năng đăng ký và gửi biểu mẫu tài khoản.",
        "pre": ["Customer chưa có tài khoản với email đang đăng ký.", "Customer truy cập được Customer Portal."],
        "post": ["Tài khoản Customer được tạo với mật khẩu đã băm.", "Email được lưu duy nhất trong hệ thống."],
        "main": [
            "Customer mở biểu mẫu đăng ký và nhập email, mật khẩu cùng các trường bắt buộc.",
            "Customer Portal gửi yêu cầu đăng ký đến Account API.",
            "Account API kiểm tra trường bắt buộc, định dạng email và chính sách mật khẩu.",
            "Account API băm mật khẩu bằng BCrypt hoặc cơ chế tương đương.",
            "Account API tạo tài khoản trong CRM Database; ràng buộc duy nhất của email được áp dụng tại cơ sở dữ liệu.",
            "Hệ thống trả kết quả tạo tài khoản thành công và Customer được hướng dẫn đăng nhập.",
        ],
        "alternative": [
            "Dữ liệu biểu mẫu không hợp lệ: API trả lỗi validation và Portal hiển thị lỗi theo từng trường.",
            "Email đã tồn tại: vi phạm ràng buộc duy nhất được bắt và trả lỗi EMAIL_ALREADY_EXISTS.",
        ],
        "exception": [
            "CRM Database không khả dụng: hệ thống trả lỗi máy chủ, không hiển thị chi tiết nội bộ.",
            "Nếu yêu cầu trùng được gửi đồng thời, chỉ một yêu cầu tạo được tài khoản; yêu cầu còn lại nhận lỗi xung đột email.",
        ],
        "rules": [
            "Email tài khoản phải duy nhất; không dựa chỉ vào bước kiểm tra sớm ở giao diện.",
            "Mật khẩu không được lưu dạng plaintext và không được trả về trong response.",
            "MVP không yêu cầu email verification; nếu bổ sung sau này phải cập nhật use case và sequence diagram.",
        ],
    },
    {
        "id": "UC-02",
        "title": "Đăng nhập, làm mới phiên và đăng xuất",
        "actors": "Customer",
        "support": "Customer Portal, Auth API, Auth Database hoặc Redis",
        "trigger": "Customer đăng nhập, phiên access token hết hạn hoặc Customer chọn đăng xuất.",
        "pre": ["Tài khoản đã tồn tại và chưa bị khóa.", "Customer có email và mật khẩu hợp lệ khi đăng nhập."],
        "post": ["Đăng nhập thành công tạo access token ngắn hạn và refresh token.", "Đăng xuất thường thu hồi phiên hiện tại và xóa refresh cookie."],
        "main": [
            "Customer gửi email và mật khẩu đến Auth API.",
            "Auth API tìm tài khoản, kiểm tra mật khẩu và trạng thái tài khoản.",
            "Auth API tạo access token ngắn hạn và refresh token ngẫu nhiên.",
            "Hệ thống chỉ lưu hash refresh token cùng userId, thời hạn, token family và trạng thái ACTIVE.",
            "Access token trả trong response body; refresh token được gửi bằng cookie HttpOnly, Secure, SameSite.",
            "Khi refresh, API khóa bản ghi token hiện tại, thu hồi token cũ với lý do ROTATED và phát token thay thế trong cùng family.",
            "Khi logout, API thu hồi chỉ refresh token hiện tại với lý do LOGOUT và xóa cookie.",
        ],
        "alternative": [
            "Sai thông tin đăng nhập hoặc tài khoản bị khóa: trả lỗi xác thực phù hợp và không tạo token.",
            "Refresh token thiếu, hết hạn hoặc bị thu hồi bởi logout/admin: trả lỗi 401 và yêu cầu đăng nhập lại.",
            "Customer có thể giữ phiên ở thiết bị khác khi logout thường; logout mọi thiết bị là thao tác riêng nếu được triển khai.",
        ],
        "exception": [
            "Refresh token đã ở trạng thái ROTATED/USED nhưng bị gửi lại: phát hiện reuse, thu hồi toàn bộ token family, ghi security event và trả REFRESH_TOKEN_REUSE_DETECTED.",
            "Kho token không khả dụng: không phát token mới và trả lỗi máy chủ an toàn.",
        ],
        "rules": [
            "Frontend chỉ giữ access token trong memory; không lưu refresh token trong localStorage hoặc nơi JavaScript đọc được.",
            "Refresh token rotation phải diễn ra trong thao tác khóa bản ghi để tránh hai lần refresh cùng sử dụng một token.",
            "Token family chỉ bị thu hồi toàn bộ khi phát hiện reuse hoặc có thao tác revoke toàn bộ thiết bị; logout thường chỉ thu hồi phiên hiện tại.",
        ],
    },
    {
        "id": "UC-03",
        "title": "Quản lý hồ sơ và sở thích",
        "actors": "Customer",
        "support": "Customer Portal, Profile API, CRM Database",
        "trigger": "Customer mở trang hồ sơ hoặc gửi thay đổi thông tin/sở thích.",
        "pre": ["Customer đã đăng nhập và có access token hợp lệ.", "Dữ liệu hồ sơ thuộc tài khoản hiện tại."],
        "post": ["Customer xem được hồ sơ hiện tại.", "Thông tin và sở thích hợp lệ được cập nhật cho đúng tài khoản."],
        "main": [
            "Customer mở trang hồ sơ; Portal gửi request kèm Authorization header.",
            "Profile API xác minh JWT trước khi truy vấn dữ liệu.",
            "API trả thông tin hồ sơ và sở thích của Customer hiện tại.",
            "Customer chỉnh sửa dữ liệu và gửi yêu cầu cập nhật.",
            "API kiểm tra định dạng, giá trị cho phép và quyền sở hữu, sau đó lưu thay đổi.",
            "Portal hiển thị hồ sơ đã cập nhật.",
        ],
        "alternative": [
            "Dữ liệu sai định dạng hoặc thiếu trường bắt buộc: trả 400 và lỗi theo trường.",
            "JWT hết hạn: trả 401 để Portal thử refresh hoặc chuyển Customer về trang đăng nhập.",
            "Tài nguyên không thuộc Customer hiện tại: trả 403 hoặc 404 theo chính sách che giấu tài nguyên.",
        ],
        "exception": ["CRM Database không khả dụng: không cập nhật một phần; trả lỗi máy chủ và ghi log không chứa secret."],
        "rules": [
            "Backend phải kiểm tra quyền sở hữu; không chỉ ẩn chức năng ở frontend.",
            "Không trả password hash hoặc các trường bảo mật trong dữ liệu hồ sơ.",
            "Dữ liệu sở thích phải dùng giá trị đã được hệ thống định nghĩa để phục vụ báo cáo Customer.",
        ],
    },
    {
        "id": "UC-04",
        "title": "Gửi feedback và rating sản phẩm sau khi mua",
        "actors": "Customer",
        "support": "Customer Portal, Feedback API, Order API, CRM Database",
        "trigger": "Customer chọn sản phẩm đã mua và gửi nội dung feedback cùng rating.",
        "pre": ["Customer đã đăng nhập.", "Customer có đơn hợp lệ chứa sản phẩm đang được đánh giá."],
        "post": ["Feedback và rating hợp lệ được lưu.", "Feedback được gắn với Customer và product duy nhất."],
        "main": [
            "Customer mở form feedback và nhập nội dung cùng rating từ 1 đến 5.",
            "Portal gửi request kèm JWT đến Feedback API.",
            "API xác minh JWT và kiểm tra order items cùng trạng thái đơn hàng của Customer có chứa sản phẩm.",
            "Nếu đơn ở CONFIRMED, SHIPPED, DELIVERED hoặc COMPLETED, API kiểm tra dữ liệu feedback.",
            "API insert feedback trực tiếp; ràng buộc duy nhất trên cặp customer/product bảo vệ dữ liệu.",
            "API trả kết quả thành công và Portal hiển thị feedback đã gửi.",
        ],
        "alternative": [
            "Customer chưa mua sản phẩm hoặc đơn chưa ở trạng thái hợp lệ: trả lỗi NOT_ELIGIBLE_TO_REVIEW.",
            "Nội dung hợp lệ nhưng Customer đã feedback sản phẩm: bắt lỗi unique constraint và trả 409 DUPLICATE_FEEDBACK.",
        ],
        "exception": [
            "Rating ngoài khoảng 1–5 hoặc nội dung không hợp lệ: trả 400 và không tạo bản ghi.",
            "Hai request feedback đồng thời: chỉ một request thành công; request còn lại nhận 409, không dùng check-trước làm cơ chế chống race condition.",
        ],
        "rules": [
            "Chỉ đơn CONFIRMED, SHIPPED, DELIVERED hoặc COMPLETED được xem là đã mua; CANCELLED không đủ điều kiện.",
            "Mỗi Customer chỉ feedback một lần cho một sản phẩm.",
            "Backend xác nhận quyền mua và quyền sở hữu đơn; không tin product/customer id do frontend tự gửi.",
        ],
    },
    {
        "id": "UC-05",
        "title": "Trả lời khảo sát",
        "actors": "Customer",
        "support": "Customer Portal, Survey API, CRM Database",
        "trigger": "Customer mở danh sách khảo sát hoặc gửi câu trả lời cho một khảo sát.",
        "pre": ["Customer đã đăng nhập.", "Khảo sát đã publish và đang trong thời gian hiệu lực."],
        "post": ["Câu trả lời hợp lệ được lưu cho đúng Customer và khảo sát.", "Customer không thể tạo response trùng cho cùng khảo sát nếu nghiệp vụ chỉ cho phép một lần."],
        "main": [
            "Customer yêu cầu danh sách khảo sát khả dụng.",
            "Survey API xác minh JWT và chỉ trả khảo sát đã publish, còn hiệu lực.",
            "Customer chọn khảo sát, nhập các câu trả lời bắt buộc và gửi form.",
            "API kiểm tra khảo sát, câu hỏi, lựa chọn và các trường bắt buộc.",
            "API lưu response; ràng buộc duy nhất và thao tác insert bảo vệ khỏi response trùng.",
            "Portal hiển thị xác nhận gửi khảo sát thành công.",
        ],
        "alternative": [
            "Khảo sát chưa publish, đã hết hiệu lực hoặc không thuộc phạm vi Customer: trả lỗi NOT_AVAILABLE.",
            "Một số câu trả lời bắt buộc bị thiếu hoặc sai kiểu: trả 400 và giữ dữ liệu form để Customer sửa.",
            "Customer đã gửi response: trả 409 DUPLICATE_RESPONSE.",
        ],
        "exception": ["Nếu lưu response thất bại, không ghi response một phần và trả lỗi máy chủ có mã theo dõi."],
        "rules": [
            "Chỉ khảo sát được phát hành mới nhận response.",
            "Các câu hỏi bắt buộc phải được kiểm tra ở backend.",
            "Một Customer chỉ gửi một response cho một khảo sát theo quy tắc MVP; nếu thay đổi phải cập nhật schema và use case.",
        ],
    },
    {
        "id": "UC-06",
        "title": "Xem, tìm kiếm và lọc sản phẩm",
        "actors": "Customer",
        "support": "Customer Portal, Catalog API, CRM Database",
        "trigger": "Customer mở danh mục hoặc nhập từ khóa/bộ lọc sản phẩm.",
        "pre": ["Danh mục có thể truy cập công khai.", "Các tham số tìm kiếm và lọc có định dạng hợp lệ nếu được gửi."],
        "post": ["Customer nhận danh sách sản phẩm phù hợp.", "Mỗi sản phẩm hiển thị đúng tên, giá, mô tả và tồn kho được phép công khai."],
        "main": [
            "Customer mở trang danh mục.",
            "Portal gửi yêu cầu lấy danh sách cùng từ khóa, bộ lọc và thông tin phân trang nếu có.",
            "Catalog API chuẩn hóa tham số, tìm kiếm theo tên/mô tả và áp dụng bộ lọc.",
            "API lấy dữ liệu sản phẩm đang được phép hiển thị.",
            "Portal hiển thị danh sách, tổng số và trạng thái rỗng nếu không có kết quả.",
        ],
        "alternative": [
            "Customer chỉ duyệt danh mục không có từ khóa: API trả danh sách mặc định theo thứ tự phù hợp.",
            "Không tìm thấy sản phẩm: trả danh sách rỗng với thông báo dễ hiểu, không xem là lỗi máy chủ.",
            "Bộ lọc không hợp lệ: bỏ bộ lọc sai hoặc trả 400 theo quy ước API thống nhất.",
        ],
        "exception": ["Database không khả dụng: trả lỗi máy chủ và không trả dữ liệu cũ gây hiểu nhầm là kết quả mới."],
        "rules": [
            "Luồng xem và tìm kiếm sản phẩm không bắt buộc Customer đăng nhập.",
            "Giá và dữ liệu dùng để đặt hàng phải được backend đọc lại khi checkout; frontend không được quyết định tổng tiền.",
            "Danh sách lớn nên hỗ trợ phân trang để tránh tải dữ liệu thừa.",
        ],
    },
    {
        "id": "UC-07",
        "title": "Giỏ hàng, đặt hàng và thanh toán payOS",
        "actors": "Customer",
        "support": "Customer Portal, Order API, CRM Database, payOS, Payment Timeout Scheduler",
        "trigger": "Customer xác nhận giỏ hàng, địa chỉ giao hàng và phương thức thanh toán.",
        "pre": ["Customer đã đăng nhập.", "Giỏ hàng có sản phẩm và số lượng hợp lệ.", "Thông tin địa chỉ và phương thức thanh toán hợp lệ."],
        "post": ["Đơn được tạo ở PENDING_PAYMENT trước khi thanh toán.", "Tồn kho được reserve an toàn hoặc được release đúng một lần.", "Chỉ webhook payOS hợp lệ mới chuyển payment PAID và order CONFIRMED."],
        "main": [
            "Customer thêm, sửa hoặc xóa sản phẩm; Portal lưu giỏ hàng qua API phía server.",
            "Customer xác nhận checkout với địa chỉ, phương thức thanh toán và promoCode nếu có.",
            "Order API bắt đầu transaction, khóa sản phẩm và đọc tồn kho hiện tại.",
            "API kiểm tra sản phẩm, số lượng và tồn kho; tổng tiền được tính từ giá hiện tại của order items.",
            "Nếu có promoCode, API kiểm tra thời hạn, điều kiện áp dụng và số lượt sử dụng, sau đó tính discountAmount và finalAmount.",
            "API tạo order PENDING_PAYMENT, lưu promo code/discount nếu hợp lệ, thêm order items và reserve tồn kho nguyên tử.",
            "API tạo orderCode duy nhất và gửi finalAmount cùng dữ liệu cần thiết đến payOS để tạo payment link.",
            "Portal chuyển Customer đến checkoutUrl. Return URL/cancelUrl chỉ hiển thị kết quả tạm thời.",
            "payOS gọi webhook; API kiểm tra signature, orderCode, số tiền finalAmount, mã dữ liệu và trạng thái.",
            "Trong transaction khóa order/payment, API ghi event chống trùng và cập nhật PAID/CONFIRMED hoặc FAILED/CANCELLED/EXPIRED.",
        ],
        "alternative": [
            "Giỏ hàng rỗng, sản phẩm không tồn tại, số lượng sai hoặc tồn kho không đủ: hủy checkout và trả 400/409.",
            "Promo code hết hạn, không đủ điều kiện hoặc hết lượt: trả lỗi mã khuyến mãi; không gửi sai finalAmount sang payOS.",
            "Customer đóng trang hoặc giao dịch thất bại: order không được xem là đã thanh toán; webhook hợp lệ có thể cập nhật FAILED/CANCELLED/EXPIRED.",
            "Webhook lặp lại: nhận diện event đã xử lý, không trừ kho hoặc ghi giao dịch lần hai, vẫn trả HTTP 2xx cho callback hợp lệ.",
        ],
        "exception": [
            "Signature hoặc dữ liệu webhook không hợp lệ: từ chối, không cập nhật order và không release/confirm sai tồn kho.",
            "Order PENDING_PAYMENT quá hạn hoặc payment link hết hạn: scheduled job khóa theo batch, chuyển EXPIRED/CANCELLED và release tồn kho đúng một lần.",
            "payOS không khả dụng sau khi order đã tạo: giữ PENDING_PAYMENT trong thời hạn xử lý và cho phép job dọn đơn treo.",
        ],
        "rules": [
            "Kiểm tra tồn kho, tạo order, thêm items và reserve stock phải nằm trong transaction với cơ chế khóa phù hợp.",
            "Payment method được lưu trong order; MVP dùng payOS, chưa yêu cầu refund production hoặc đối soát production.",
            "Return URL/cancelUrl chỉ phục vụ hiển thị; webhook là nguồn xác nhận thanh toán chính thức.",
            "Webhook phải kiểm tra signature, orderCode, amount bằng finalAmount và idempotency; callback hợp lệ đã xử lý vẫn trả 2xx.",
        ],
    },
    {
        "id": "UC-08",
        "title": "Xem lịch sử và trạng thái đơn hàng",
        "actors": "Customer",
        "support": "Customer Portal, Order API, CRM Database",
        "trigger": "Customer mở lịch sử đơn hoặc chọn một đơn để xem chi tiết.",
        "pre": ["Customer đã đăng nhập với JWT hợp lệ.", "Đơn hàng tồn tại và thuộc tài khoản Customer hiện tại."],
        "post": ["Customer nhận danh sách hoặc chi tiết đơn của chính mình.", "Trạng thái đơn và thanh toán được hiển thị theo dữ liệu backend mới nhất."],
        "main": [
            "Customer mở trang lịch sử đơn hàng.",
            "Portal gửi request kèm Authorization header.",
            "Order API xác minh JWT và truy vấn các order có customerId đúng tài khoản hiện tại.",
            "API trả danh sách có phân trang, trạng thái đơn, trạng thái thanh toán và thời gian tạo.",
            "Customer chọn một order; API kiểm tra quyền sở hữu rồi trả order items, tổng tiền và trạng thái order/payment hiện tại.",
            "Portal hiển thị trạng thái dễ hiểu cho Customer.",
        ],
        "alternative": [
            "Customer chưa có đơn: trả danh sách rỗng và trạng thái hướng dẫn mua hàng.",
            "Order không thuộc Customer hiện tại: trả 403 hoặc 404 theo chính sách, không tiết lộ dữ liệu order.",
            "JWT hết hạn: trả 401 để Portal refresh hoặc yêu cầu đăng nhập lại.",
        ],
        "exception": ["Database không khả dụng: trả lỗi máy chủ và không hiển thị trạng thái cũ như dữ liệu hiện tại."],
        "rules": [
            "Customer chỉ được xem order của chính mình; backend phải lọc theo danh tính đã xác thực.",
            "PENDING_PAYMENT không được hiển thị là PAID nếu chưa có webhook payOS hợp lệ.",
            "Order status chỉ chuyển theo luồng nghiệp vụ đã thống nhất và không được sửa trực tiếp bởi Customer.",
        ],
    },
]


def set_run_font(run, name="Aptos", size=10.5, bold=False, color="000000", italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D9D9D9", size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn("w:" + key))
        if node is None:
            node = OxmlElement("w:" + key)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def format_table(table, header_fill="1F4E78", font_size=8.5):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    mark_header_row(table.rows[0])
    for r_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            set_cell_margins(cell)
            if r_index == 0:
                set_cell_shading(cell, header_fill)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    set_run_font(run, size=font_size, bold=(r_index == 0), color=("FFFFFF" if r_index == 0 else "000000"))


def add_paragraph(doc, text="", style=None, bold_prefix=None, italic=False):
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.12
    if bold_prefix and text.startswith(bold_prefix):
        run = paragraph.add_run(bold_prefix)
        set_run_font(run, bold=True)
        run = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(run, italic=italic)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, italic=italic)
    return paragraph


def add_bullets(doc, items, style="List Bullet"):
    for item in items:
        paragraph = doc.add_paragraph(style=style)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.08
        run = paragraph.add_run(item)
        set_run_font(run)


def add_numbered(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.08
        run = paragraph.add_run(item)
        set_run_font(run)


def add_key_value_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Thuộc tính"
    table.rows[0].cells[1].text = "Nội dung"
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value
    format_table(table, font_size=9)
    for row in table.rows[1:]:
        set_cell_shading(row.cells[0], "EAF2F8")
        for run in row.cells[0].paragraphs[0].runs:
            set_run_font(run, size=9, bold=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_section_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    set_run_font(run, size=(15 if level == 1 else 11.5), bold=True)
    return paragraph


def configure_styles(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)

    for style_name, size in (("Heading 1", 15), ("Heading 2", 11.5), ("Heading 3", 10.5)):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)

    title = doc.styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Part III | Customer CRM and Sales Use Case Specifications")
    set_run_font(footer_run, size=8.5, color="666666")


def build_document():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("Đặc tả Use Case Customer CRM và Sales")
    set_run_font(run, name="Aptos Display", size=22, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run("Part III | Phiên bản MVP")
    set_run_font(run, size=12, bold=True, color="1F4E78")

    add_key_value_table(doc, [
        ("Mục đích", "Mô tả thống nhất các luồng Customer CRM và Sales để làm cơ sở cho thiết kế, triển khai và kiểm thử."),
        ("Phạm vi", "Tám chức năng Customer: tài khoản, xác thực, hồ sơ, feedback, khảo sát, catalog, giỏ hàng/đặt hàng và lịch sử đơn."),
        ("Tài liệu liên quan", "PRD hệ thống CRM và Quản lý bán hàng; Customer Use Case Diagram; các sequence diagram III-01 đến III-08."),
        ("Thanh toán MVP", "payOS; Return URL/cancelUrl chỉ hiển thị kết quả, webhook hợp lệ là nguồn xác nhận thanh toán."),
        ("Ngoài phạm vi", "Email verification, thanh toán production/refund production, đối soát production và tích hợp đơn vị vận chuyển thực tế."),
    ])

    add_section_heading(doc, "1. Giới thiệu", 1)
    add_paragraph(doc, "Tài liệu này đặc tả các Use Case của Customer Portal trong hệ thống CRM và quản lý bán hàng. Mỗi đặc tả mô tả actor, điều kiện, luồng xử lý và quy tắc nghiệp vụ ở mức đủ chi tiết để nhóm thống nhất trước khi triển khai và kiểm thử.")
    add_paragraph(doc, "Customer chỉ được truy cập dữ liệu của chính mình. Backend chịu trách nhiệm xác thực, phân quyền, kiểm tra dữ liệu và bảo toàn transaction; frontend chỉ hỗ trợ nhập liệu, điều hướng và hiển thị trạng thái.")

    add_section_heading(doc, "2. Bảng truy vết", 1)
    trace = doc.add_table(rows=1, cols=5)
    trace.style = "Table Grid"
    headers = ["Use Case", "Chức năng", "FR / Issue", "Diagram", "Sequence"]
    for cell, value in zip(trace.rows[0].cells, headers):
        cell.text = value
    for row in TRACEABILITY:
        cells = trace.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value
    format_table(trace, font_size=8.2)
    add_paragraph(doc, "Các mã FR và Issue trong bảng giúp đối chiếu trực tiếp giữa yêu cầu chức năng, diagram tổng hợp và sequence diagram chi tiết.", italic=True)

    for index, uc in enumerate(USE_CASES):
        doc.add_page_break()
        add_section_heading(doc, f"{index + 3}. {uc['id']} - {uc['title']}", 1)
        add_key_value_table(doc, [
            ("Actor chính", uc["actors"]),
            ("Actor hỗ trợ", uc["support"]),
            ("Kích hoạt", uc["trigger"]),
            ("Tiền điều kiện", " ".join(uc["pre"])),
            ("Hậu điều kiện", " ".join(uc["post"])),
        ])

        add_section_heading(doc, "Luồng chính", 2)
        add_numbered(doc, uc["main"])
        add_section_heading(doc, "Luồng thay thế", 2)
        add_bullets(doc, uc["alternative"])
        add_section_heading(doc, "Luồng ngoại lệ", 2)
        add_bullets(doc, uc["exception"])
        add_section_heading(doc, "Quy tắc nghiệp vụ và validation", 2)
        add_bullets(doc, uc["rules"])

    doc.add_page_break()
    add_section_heading(doc, "11. Ghi chú nhất quán giữa các Use Case", 1)
    add_bullets(doc, [
        "Access token có thời hạn ngắn và chỉ nằm trong memory của frontend; refresh token được lưu dạng hash và gửi qua HttpOnly, Secure, SameSite cookie.",
        "Customer chỉ feedback sản phẩm sau khi có order hợp lệ chứa sản phẩm; unique constraint xử lý trường hợp feedback đồng thời.",
        "Order phải tính lại từ dữ liệu backend, áp dụng promo code trước khi gửi finalAmount sang payOS và không xem Return URL là bằng chứng thanh toán.",
        "Webhook payOS phải xác minh signature, orderCode, amount và tính idempotent; order PENDING_PAYMENT quá hạn phải được job dọn đúng một lần.",
        "Mọi thay đổi liên quan schema, trạng thái đơn hoặc quy tắc nghiệp vụ phải cập nhật ERD/data dictionary và sequence tương ứng trước khi merge.",
    ])
    add_paragraph(doc, "Tài liệu này là deliverable của Issue #53 và được thiết kế để nhất quán với Issue #52 cùng các sequence diagram Customer trong thư mục docs/diagrams/customer.", italic=True)

    doc.core_properties.title = "Đặc tả Use Case Customer CRM và Sales"
    doc.core_properties.subject = "Part III Customer CRM and Sales"
    doc.core_properties.author = "Project Team"
    doc.core_properties.comments = "Customer use case specifications for MVP"
    doc.save(OUTPUT)


if __name__ == "__main__":
    build_document()
    print(OUTPUT)
