from pathlib import Path
import importlib.util

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "use-cases" / "admin" / "admin-use-case-specifications.docx"
CUSTOMER_BUILDER = ROOT / "tools" / "build_customer_usecase_specifications.py"


def load_customer_helpers():
    spec = importlib.util.spec_from_file_location("customer_builder", CUSTOMER_BUILDER)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


ADMIN_TRACEABILITY = [
    ("UC-A01", "Tiếp nhận và xử lý phản hồi khách hàng", "Admin-01", "01", "01-feedback-handling.puml"),
    ("UC-A02", "Xem thông tin khách hàng", "Admin-02", "02", "02-view-customer-information.puml"),
    ("UC-A04", "Khóa và mở tài khoản khách hàng", "Admin-04", "04", "04-lock-unlock-account.puml"),
    ("UC-A05", "Báo cáo khách hàng", "Admin-05", "05", "05-customer-report.puml"),
    ("UC-A06", "Quản lý danh mục và thông tin sản phẩm", "Admin-06", "06", "06-product-management.puml"),
    ("UC-A07", "Tạo và gửi bảng khảo sát", "Admin-07", "07", "07-survey-create-send.puml"),
    ("UC-A08", "Thống kê và phân tích kết quả khảo sát", "Admin-08", "08", "08-survey-analytics.puml"),
    ("UC-A09", "Quản lý và duyệt đơn hàng bán", "Admin-09", "09", "09-sales-order-approval.puml"),
    ("UC-A10", "Thêm khách hàng", "Admin-10", "10", "10-add-customer.puml"),
    ("UC-A11", "Xóa khách hàng", "Admin-11", "11", "11-delete-customer.puml"),
    ("UC-A12", "Thống kê và báo cáo doanh thu", "Admin-12 / #14", "12", "12-sales-revenue-reporting.puml"),
]


ADMIN_USE_CASES = [
    {
        "id": "UC-A01",
        "title": "Tiếp nhận và xử lý phản hồi khách hàng",
        "actors": "Admin",
        "support": "Admin Portal, Feedback API, CRM Database, Notification Service",
        "trigger": "Admin mở danh sách phản hồi và chọn một phản hồi cần xử lý.",
        "pre": ["Admin đã đăng nhập và có quyền quản lý phản hồi.", "Phản hồi tồn tại trong hệ thống và đang chờ xử lý."],
        "post": ["Phản hồi được chuyển sang trạng thái đã duyệt hoặc bị từ chối.", "Customer nhận được kết quả xử lý theo kênh thông báo được cấu hình."],
        "main": [
            "Admin mở trang quản lý phản hồi.",
            "Admin Portal gửi yêu cầu lấy danh sách phản hồi đang chờ xử lý.",
            "Feedback API xác minh quyền Admin và trả danh sách phản hồi phù hợp.",
            "Admin chọn một phản hồi và chọn Duyệt hoặc Từ chối.",
            "API kiểm tra trạng thái hiện tại, dữ liệu xử lý và lý do từ chối nếu có.",
            "API cập nhật trạng thái, ghi nhận người xử lý và tạo thông báo cho Customer.",
            "Portal hiển thị kết quả xử lý cho Admin.",
        ],
        "alternative": [
            "Không có phản hồi chờ xử lý: hiển thị danh sách rỗng và hướng dẫn phù hợp.",
            "Admin chỉ xem được phản hồi nhưng không có quyền xử lý: trả lỗi 403 và không thay đổi dữ liệu.",
            "Admin từ chối phản hồi: bắt buộc nhập lý do trước khi cập nhật trạng thái.",
        ],
        "exception": [
            "Phản hồi đã được xử lý bởi Admin khác: trả thông báo trạng thái mới nhất và không ghi đè kết quả.",
            "Notification Service không khả dụng: vẫn giữ kết quả xử lý, ghi nhận việc gửi lại và không tạo bản ghi xử lý trùng.",
        ],
        "rules": [
            "Chỉ Admin có quyền phù hợp mới được duyệt hoặc từ chối phản hồi.",
            "Một phản hồi chỉ được chuyển trạng thái xử lý một lần theo quy tắc nghiệp vụ.",
            "Lý do từ chối phải được lưu để phục vụ tra cứu và khiếu nại.",
        ],
    },
    {
        "id": "UC-A02",
        "title": "Xem thông tin khách hàng",
        "actors": "Admin",
        "support": "Admin Portal, Customer API, CRM Database",
        "trigger": "Admin tìm kiếm hoặc chọn một Customer trong danh sách quản lý.",
        "pre": ["Admin đã đăng nhập và có quyền xem Customer.", "Customer có thể được tra cứu theo mã, email hoặc thông tin được cho phép."],
        "post": ["Admin xem được thông tin Customer trong phạm vi quyền.", "Dữ liệu bảo mật không cần thiết không bị hiển thị."],
        "main": [
            "Admin nhập tiêu chí tìm kiếm hoặc mở danh sách Customer.",
            "Portal gửi yêu cầu đến Customer API kèm thông tin xác thực Admin.",
            "API xác minh quyền và chuẩn hóa tiêu chí tìm kiếm.",
            "API truy vấn danh sách hoặc chi tiết Customer từ CRM Database.",
            "API loại bỏ các trường nhạy cảm không thuộc phạm vi hiển thị.",
            "Portal hiển thị thông tin Customer và trạng thái tài khoản.",
        ],
        "alternative": [
            "Không tìm thấy Customer: trả danh sách rỗng hoặc 404 cho màn hình chi tiết.",
            "Admin nhập tiêu chí không hợp lệ: trả lỗi theo trường để Admin sửa lại.",
            "Admin không đủ quyền: trả 403 và không tiết lộ dữ liệu Customer.",
        ],
        "exception": [
            "CRM Database không khả dụng: trả lỗi máy chủ và ghi log kỹ thuật không chứa dữ liệu nhạy cảm.",
            "Dữ liệu Customer thiếu trường không bắt buộc: hiển thị trạng thái chưa cập nhật thay vì làm hỏng toàn bộ danh sách.",
        ],
        "rules": [
            "Mọi request phải kiểm tra quyền ở backend; không chỉ ẩn nút trên giao diện.",
            "Không hiển thị password hash, refresh token hoặc secret cho Admin.",
            "Kết quả danh sách cần hỗ trợ phân trang khi số lượng Customer lớn.",
        ],
    },
    {
        "id": "UC-A04",
        "title": "Khóa và mở tài khoản khách hàng",
        "actors": "Admin",
        "support": "Admin Portal, Customer API, CRM Database, Audit Log",
        "trigger": "Admin chọn thao tác khóa hoặc mở khóa trên tài khoản Customer.",
        "pre": ["Admin đã đăng nhập và có quyền quản lý trạng thái tài khoản.", "Tài khoản Customer tồn tại và trạng thái hiện tại được xác định."],
        "post": ["Trạng thái tài khoản được cập nhật đúng một lần.", "Thao tác được ghi audit để truy vết người thực hiện và thời điểm."],
        "main": [
            "Admin tìm và chọn tài khoản Customer.",
            "Portal gửi yêu cầu khóa hoặc mở khóa đến Customer API.",
            "API xác minh quyền Admin và đọc trạng thái hiện tại của tài khoản.",
            "API kiểm tra thao tác có phù hợp với trạng thái hiện tại hay không.",
            "API cập nhật trạng thái tài khoản và ghi audit log.",
            "Portal hiển thị trạng thái mới.",
        ],
        "alternative": [
            "Tài khoản đã ở trạng thái yêu cầu: trả kết quả idempotent và không tạo audit event trùng.",
            "Admin không đủ quyền: trả 403 và không cập nhật trạng thái.",
            "Customer không tồn tại: trả 404 mà không tiết lộ thông tin tài khoản khác.",
        ],
        "exception": [
            "Hai Admin thao tác đồng thời: chỉ một cập nhật hợp lệ theo trạng thái mới nhất, thao tác còn lại phải đọc lại kết quả.",
            "Lưu audit thất bại: không hoàn tất thay đổi trạng thái nếu hệ thống yêu cầu audit bắt buộc.",
        ],
        "rules": [
            "Tài khoản bị khóa không được đăng nhập hoặc sử dụng phiên mới.",
            "Thao tác mở khóa không tự động thay đổi dữ liệu hồ sơ Customer.",
            "Mọi chuyển đổi trạng thái phải có người thực hiện và lý do theo chính sách audit.",
        ],
    },
    {
        "id": "UC-A05",
        "title": "Báo cáo khách hàng",
        "actors": "Admin",
        "support": "Admin Portal, Customer Report API, CRM Database",
        "trigger": "Admin mở chức năng báo cáo và chọn bộ lọc nhân khẩu học, độ tuổi hoặc sở thích.",
        "pre": ["Admin có quyền xem báo cáo Customer.", "Dữ liệu Customer và các giá trị phân loại đã được chuẩn hóa."],
        "post": ["Báo cáo được tổng hợp theo bộ lọc đã chọn.", "Số liệu hiển thị có thể được dùng cho phân tích và quyết định kinh doanh."],
        "main": [
            "Admin chọn loại báo cáo và khoảng thời gian hoặc bộ lọc.",
            "Portal gửi yêu cầu đến Customer Report API.",
            "API xác minh quyền và kiểm tra tham số báo cáo.",
            "API tổng hợp số lượng Customer theo nhóm nhân khẩu học, độ tuổi và sở thích.",
            "API trả dữ liệu tổng hợp cho Portal.",
            "Portal hiển thị bảng hoặc biểu đồ báo cáo.",
        ],
        "alternative": [
            "Admin không chọn bộ lọc: hệ thống dùng phạm vi mặc định đã công bố.",
            "Không có dữ liệu phù hợp: hiển thị báo cáo rỗng và thông tin giải thích.",
            "Tham số ngoài phạm vi cho phép: trả lỗi 400 và không chạy truy vấn báo cáo.",
        ],
        "exception": [
            "Dữ liệu phân loại chưa đầy đủ: nhóm các giá trị thiếu vào trạng thái chưa xác định và ghi nhận trong báo cáo.",
            "Báo cáo quá lớn: yêu cầu phân trang hoặc xử lý bất đồng bộ theo quy định của hệ thống.",
        ],
        "rules": [
            "Báo cáo chỉ trả dữ liệu tổng hợp, không làm lộ thông tin Customer không cần thiết.",
            "Các bộ lọc phải dùng cùng định nghĩa với dữ liệu hồ sơ Customer.",
            "Thời điểm tạo và bộ lọc của báo cáo cần được hiển thị để người xem hiểu phạm vi số liệu.",
        ],
    },
    {
        "id": "UC-A06",
        "title": "Quản lý danh mục và thông tin sản phẩm",
        "actors": "Admin",
        "support": "Admin Portal, Product API, CRM Database",
        "trigger": "Admin tạo, sửa, phân loại hoặc ngừng kinh doanh một sản phẩm.",
        "pre": ["Admin có quyền quản lý sản phẩm.", "Dữ liệu sản phẩm hoặc danh mục được chọn tồn tại khi thao tác cập nhật."],
        "post": ["Thông tin sản phẩm và danh mục hợp lệ được lưu.", "Catalog của Customer chỉ hiển thị dữ liệu đã được công bố."],
        "main": [
            "Admin mở danh sách sản phẩm và danh mục.",
            "Admin chọn tạo mới hoặc chỉnh sửa thông tin sản phẩm.",
            "Portal gửi dữ liệu đến Product API.",
            "API kiểm tra trường bắt buộc, giá, danh mục và trạng thái công bố.",
            "API kiểm tra trùng mã hoặc tên theo quy tắc nghiệp vụ.",
            "API lưu thay đổi và trả dữ liệu sản phẩm mới nhất.",
            "Portal hiển thị kết quả cập nhật.",
        ],
        "alternative": [
            "Dữ liệu sai định dạng hoặc giá không hợp lệ: trả lỗi theo trường.",
            "Mã sản phẩm đã tồn tại: từ chối tạo mới và yêu cầu Admin dùng mã khác.",
            "Admin ngừng sản phẩm: chuyển sang trạng thái không hiển thị thay vì làm mất lịch sử đơn hàng.",
        ],
        "exception": [
            "Sản phẩm đã có trong đơn hàng: không xóa vật lý nếu việc đó làm mất lịch sử; áp dụng chính sách ngừng kinh doanh.",
            "Cập nhật đồng thời: API kiểm tra phiên bản hoặc đọc lại dữ liệu trước khi ghi để tránh ghi đè ngoài ý muốn.",
        ],
        "rules": [
            "Giá được lưu và kiểm tra ở backend; Customer không được tự quyết định giá checkout.",
            "Sản phẩm chưa công bố không xuất hiện trong Catalog công khai.",
            "Thay đổi ảnh, mô tả hoặc danh mục phải giữ được liên kết với sản phẩm và lịch sử đơn hàng.",
        ],
    },
    {
        "id": "UC-A07",
        "title": "Tạo và gửi bảng khảo sát",
        "actors": "Admin",
        "support": "Admin Portal, Survey API, CRM Database, Notification Service",
        "trigger": "Admin tạo khảo sát mới, hoàn thiện câu hỏi và chọn phát hành cho Customer.",
        "pre": ["Admin có quyền quản lý khảo sát.", "Khảo sát có tiêu đề, câu hỏi và cấu hình thời gian hợp lệ trước khi phát hành."],
        "post": ["Khảo sát được lưu với trạng thái phù hợp.", "Customer thuộc phạm vi nhận được thông tin khảo sát khi khảo sát được phát hành."],
        "main": [
            "Admin mở chức năng tạo khảo sát.",
            "Admin nhập tiêu đề, thời gian hiệu lực, câu hỏi và các lựa chọn trả lời.",
            "Portal gửi dữ liệu đến Survey API.",
            "API kiểm tra cấu trúc, câu hỏi bắt buộc và thời gian khảo sát.",
            "API lưu khảo sát ở trạng thái bản nháp.",
            "Admin xem lại và chọn phát hành.",
            "API chuyển khảo sát sang PUBLISHED và tạo yêu cầu thông báo cho Customer mục tiêu.",
        ],
        "alternative": [
            "Thiếu câu hỏi hoặc cấu hình không hợp lệ: giữ bản nháp và trả lỗi để Admin chỉnh sửa.",
            "Admin lưu bản nháp nhưng chưa phát hành: Customer chưa thể xem hoặc trả lời.",
            "Admin chỉnh sửa khảo sát đã phát hành: chỉ cho phép thay đổi theo chính sách, tránh làm sai response đã nhận.",
        ],
        "exception": [
            "Gửi thông báo thất bại: khảo sát vẫn giữ trạng thái đã phát hành, ghi nhận danh sách cần gửi lại.",
            "Hai yêu cầu phát hành đồng thời: chỉ một yêu cầu chuyển trạng thái và tạo thông báo.",
        ],
        "rules": [
            "Chỉ khảo sát PUBLISHED trong thời gian hiệu lực mới được Customer nhìn thấy.",
            "Câu hỏi và lựa chọn đã dùng trong response không được thay đổi tùy tiện.",
            "Việc phát hành phải có người thực hiện và thời điểm để phục vụ audit.",
        ],
    },
    {
        "id": "UC-A08",
        "title": "Thống kê và phân tích kết quả khảo sát",
        "actors": "Admin",
        "support": "Admin Portal, Survey Analytics API, CRM Database",
        "trigger": "Admin chọn một khảo sát đã có response và mở trang phân tích.",
        "pre": ["Admin có quyền xem thống kê khảo sát.", "Khảo sát tồn tại và dữ liệu response đã được lưu hợp lệ."],
        "post": ["Admin xem được số lượng, tỷ lệ và phân bố câu trả lời.", "Các giá trị không đủ dữ liệu được thể hiện rõ thay vì suy diễn."],
        "main": [
            "Admin chọn khảo sát và phạm vi phân tích.",
            "Portal gửi yêu cầu đến Survey Analytics API.",
            "API xác minh quyền và kiểm tra khảo sát.",
            "API tổng hợp response theo câu hỏi, lựa chọn và các nhóm được cho phép.",
            "API trả dữ liệu tổng hợp cùng số lượng mẫu.",
            "Portal hiển thị bảng, biểu đồ và trạng thái dữ liệu.",
        ],
        "alternative": [
            "Khảo sát chưa có response: hiển thị trạng thái chưa có dữ liệu.",
            "Admin lọc theo khoảng thời gian: API chỉ tổng hợp response thuộc phạm vi đó.",
            "Một câu hỏi không có câu trả lời hợp lệ: hiển thị tỷ lệ thiếu và không đưa vào mẫu hợp lệ.",
        ],
        "exception": [
            "Dữ liệu response không nhất quán: loại bản ghi lỗi khỏi thống kê, ghi log và báo rõ số lượng bị loại.",
            "Tập dữ liệu lớn: thực hiện phân trang hoặc job tổng hợp theo khả năng hệ thống.",
        ],
        "rules": [
            "Báo cáo chỉ hiển thị dữ liệu tổng hợp, không lộ nội dung nhận diện Customer nếu không có quyền riêng.",
            "Luôn hiển thị cỡ mẫu và phạm vi thời gian của thống kê.",
            "Cách tính tỷ lệ phải thống nhất giữa báo cáo và dashboard.",
        ],
    },
    {
        "id": "UC-A09",
        "title": "Quản lý và duyệt đơn hàng bán",
        "actors": "Admin",
        "support": "Admin Portal, Order API, CRM Database, Notification Service",
        "trigger": "Admin mở danh sách đơn hàng và chọn đơn cần kiểm tra hoặc duyệt.",
        "pre": ["Admin có quyền quản lý đơn hàng.", "Đơn hàng tồn tại và trạng thái hiện tại được xác định."],
        "post": ["Đơn hàng chuyển sang trạng thái phù hợp sau khi duyệt hoặc từ chối.", "Customer nhận được trạng thái mới và lịch sử xử lý được ghi nhận."],
        "main": [
            "Admin mở danh sách đơn hàng theo trạng thái hoặc thời gian.",
            "Portal gửi request đến Order API.",
            "API trả thông tin đơn, Customer, items, tổng tiền và trạng thái thanh toán.",
            "Admin mở chi tiết và kiểm tra thông tin đơn.",
            "Admin chọn duyệt, từ chối hoặc thao tác xử lý được phép.",
            "API kiểm tra chuyển trạng thái hợp lệ và cập nhật đơn trong transaction.",
            "API tạo thông báo cho Customer và Portal hiển thị kết quả.",
        ],
        "alternative": [
            "Đơn đã được Admin khác xử lý: hiển thị trạng thái mới nhất và không ghi đè.",
            "Đơn không đủ điều kiện duyệt: yêu cầu Admin nhập lý do từ chối hoặc chuyển xử lý.",
            "Admin chỉ có quyền xem: không hiển thị thao tác thay đổi trạng thái.",
        ],
        "exception": [
            "Trạng thái thanh toán chưa được webhook xác nhận: không chuyển đơn thành đã thanh toán chỉ dựa vào Return URL.",
            "Notification Service không khả dụng: lưu trạng thái đơn, tạo nhiệm vụ gửi lại và không xử lý đơn lần hai.",
        ],
        "rules": [
            "Chỉ chuyển trạng thái theo state machine đã thống nhất.",
            "Mọi kiểm tra quyền, tồn kho và thanh toán phải do backend thực hiện.",
            "Không được sửa tổng tiền đã xác nhận thanh toán nếu không có quy trình điều chỉnh riêng.",
        ],
    },
    {
        "id": "UC-A10",
        "title": "Thêm khách hàng",
        "actors": "Admin",
        "support": "Admin Portal, Customer API, CRM Database",
        "trigger": "Admin chọn thêm Customer và gửi biểu mẫu thông tin.",
        "pre": ["Admin có quyền tạo Customer.", "Các trường bắt buộc đã được nhập theo định dạng cho phép."],
        "post": ["Customer mới được tạo với thông tin hợp lệ.", "Email hoặc mã định danh được bảo vệ bởi ràng buộc duy nhất phù hợp."],
        "main": [
            "Admin mở biểu mẫu thêm Customer.",
            "Admin nhập thông tin tài khoản và hồ sơ cơ bản.",
            "Portal gửi yêu cầu đến Customer API.",
            "API kiểm tra dữ liệu bắt buộc, định dạng và quyền tạo.",
            "API tạo Customer và lưu mật khẩu theo cơ chế băm nếu quy trình có cấp tài khoản.",
            "Portal hiển thị kết quả tạo thành công.",
        ],
        "alternative": [
            "Thiếu hoặc sai dữ liệu: trả lỗi theo trường và không tạo bản ghi một phần.",
            "Email đã tồn tại: bắt lỗi ràng buộc duy nhất và yêu cầu Admin kiểm tra Customer hiện có.",
            "Admin hủy biểu mẫu: không phát sinh thay đổi.",
        ],
        "exception": [
            "Hai yêu cầu tạo cùng email: chỉ một yêu cầu thành công, yêu cầu còn lại nhận lỗi xung đột.",
            "CRM Database không khả dụng: trả lỗi máy chủ và không hiển thị chi tiết nội bộ.",
        ],
        "rules": [
            "Không dựa chỉ vào bước kiểm tra trùng ở giao diện; cơ sở dữ liệu phải bảo vệ duy nhất.",
            "Mật khẩu không được lưu plaintext hoặc hiển thị lại sau khi tạo.",
            "Thao tác tạo Customer cần ghi người thực hiện nếu hệ thống bật audit.",
        ],
    },
    {
        "id": "UC-A11",
        "title": "Xóa khách hàng",
        "actors": "Admin",
        "support": "Admin Portal, Customer API, CRM Database, Audit Log",
        "trigger": "Admin chọn Customer và yêu cầu xóa hoặc vô hiệu hóa tài khoản.",
        "pre": ["Admin có quyền quản lý vòng đời Customer.", "Customer tồn tại và không có thao tác xóa khác đang xử lý."],
        "post": ["Customer được xóa hoặc chuyển sang trạng thái không hoạt động theo chính sách lưu trữ.", "Dữ liệu lịch sử cần thiết cho đơn hàng và audit vẫn được bảo toàn."],
        "main": [
            "Admin tìm và chọn Customer cần xóa.",
            "Portal hiển thị cảnh báo và yêu cầu Admin xác nhận.",
            "API kiểm tra quyền, trạng thái Customer và các liên kết nghiệp vụ.",
            "API áp dụng chính sách xóa hoặc vô hiệu hóa phù hợp.",
            "API ghi audit log và trả kết quả.",
            "Portal cập nhật danh sách Customer.",
        ],
        "alternative": [
            "Admin hủy xác nhận: không thay đổi dữ liệu.",
            "Customer có đơn hàng hoặc dữ liệu cần lưu: dùng vô hiệu hóa/ẩn tài khoản thay vì xóa vật lý.",
            "Customer không tồn tại: trả 404 và không thực hiện thao tác khác.",
        ],
        "exception": [
            "Có thao tác cập nhật đồng thời: từ chối hoặc yêu cầu đọc lại trạng thái trước khi xóa.",
            "Audit log không ghi được: không hoàn tất thao tác nếu audit là bắt buộc.",
        ],
        "rules": [
            "Không xóa dữ liệu làm mất lịch sử đơn hàng, thanh toán, feedback hoặc audit cần thiết.",
            "Sau khi vô hiệu hóa, Customer không được đăng nhập hoặc tạo giao dịch mới.",
            "Thao tác xóa phải có xác nhận rõ ràng và ghi nhận Admin thực hiện.",
        ],
    },
    {
        "id": "UC-A12",
        "title": "Thống kê và báo cáo doanh thu",
        "actors": "Admin",
        "support": "Admin Portal, Revenue Report API, CRM Database",
        "trigger": "Admin chọn khoảng thời gian và mở báo cáo doanh thu.",
        "pre": ["Admin có quyền xem báo cáo doanh thu.", "Đơn hàng và trạng thái thanh toán đã được cập nhật theo nguồn dữ liệu chính thức."],
        "post": ["Admin xem được doanh thu theo thời gian, sản phẩm bán chạy và lợi nhuận theo phạm vi đã chọn.", "Báo cáo thể hiện rõ cách tính và phạm vi dữ liệu."],
        "main": [
            "Admin chọn khoảng thời gian và loại báo cáo.",
            "Portal gửi yêu cầu đến Revenue Report API.",
            "API xác minh quyền và kiểm tra khoảng thời gian.",
            "API tổng hợp các đơn hợp lệ theo trạng thái thanh toán và trạng thái đơn.",
            "API tính doanh thu, sản phẩm bán chạy và lợi nhuận theo định nghĩa báo cáo.",
            "Portal hiển thị bảng hoặc biểu đồ kết quả.",
        ],
        "alternative": [
            "Admin không chọn khoảng thời gian: dùng khoảng mặc định được cấu hình.",
            "Không có đơn hợp lệ: hiển thị số liệu bằng 0 và trạng thái không có dữ liệu.",
            "Admin chọn khoảng thời gian quá lớn: yêu cầu thu hẹp hoặc chạy báo cáo bất đồng bộ.",
        ],
        "exception": [
            "Dữ liệu đơn hoặc thanh toán chưa nhất quán: loại khỏi báo cáo chính, đánh dấu cần đối soát và ghi log.",
            "Database không khả dụng: trả lỗi máy chủ, không hiển thị báo cáo cũ như số liệu hiện tại.",
        ],
        "rules": [
            "Chỉ tính các giao dịch đã được xác nhận theo trạng thái thanh toán chính thức.",
            "Khoảng thời gian, múi giờ và công thức phải được hiển thị cùng báo cáo.",
            "Báo cáo phải phân biệt doanh thu, giảm giá, phí và lợi nhuận theo định nghĩa nghiệp vụ.",
        ],
    },
]


def build_document():
    helper = load_customer_helpers()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    helper.configure_styles(doc)

    footer = doc.sections[0].footer.paragraphs[0]
    if footer.runs:
        footer.runs[0].text = "Part III | Admin CRM and Sales Use Case Specifications"

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("Đặc tả Use Case Admin CRM và Sales")
    helper.set_run_font(run, name="Aptos Display", size=22, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run("Part III | Phiên bản MVP")
    helper.set_run_font(run, size=12, bold=True, color="1F4E78")

    helper.add_key_value_table(doc, [
        ("Mục đích", "Mô tả thống nhất các luồng quản trị Admin để làm cơ sở cho thiết kế, triển khai, phân quyền và kiểm thử."),
        ("Phạm vi", "Mười một chức năng Admin: phản hồi, Customer, tài khoản, báo cáo, sản phẩm, khảo sát, đơn hàng và doanh thu."),
        ("Tài liệu liên quan", "PRD hệ thống CRM và Quản lý bán hàng; Admin Use Case Diagram; các sequence diagram Admin 01, 02, 04 đến 12."),
        ("Phân quyền", "Mọi thao tác thay đổi phải được kiểm tra ở backend và ghi audit theo chính sách hệ thống."),
        ("Ngoài phạm vi", "Các chức năng vận hành production chưa có yêu cầu riêng như đối soát, hoàn tiền và tích hợp vận chuyển thực tế."),
    ])

    helper.add_section_heading(doc, "1. Giới thiệu", 1)
    helper.add_paragraph(doc, "Tài liệu này đặc tả các Use Case của Admin Portal trong hệ thống CRM và quản lý bán hàng. Mỗi đặc tả mô tả actor, điều kiện, luồng xử lý và quy tắc nghiệp vụ ở mức đủ chi tiết để nhóm thống nhất trước khi triển khai và kiểm thử.")
    helper.add_paragraph(doc, "Admin chỉ được thực hiện chức năng trong phạm vi quyền được cấp. Backend chịu trách nhiệm xác thực, phân quyền, kiểm tra dữ liệu, bảo toàn transaction và ghi audit; frontend chỉ hỗ trợ nhập liệu, điều hướng và hiển thị trạng thái.")

    helper.add_section_heading(doc, "2. Bảng truy vết", 1)
    trace = doc.add_table(rows=1, cols=5)
    trace.style = "Table Grid"
    for cell, value in zip(trace.rows[0].cells, ["Use Case", "Chức năng", "FR / Task", "Diagram", "Sequence"]):
        cell.text = value
    for row in ADMIN_TRACEABILITY:
        cells = trace.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value
    helper.format_table(trace, font_size=8.2)
    helper.add_paragraph(doc, "Các mã Admin trong bảng giúp đối chiếu giữa yêu cầu quản trị, diagram tổng hợp và sequence diagram chi tiết.", italic=True)

    for index, uc in enumerate(ADMIN_USE_CASES):
        doc.add_page_break()
        helper.add_section_heading(doc, f"{index + 3}. {uc['id']} - {uc['title']}", 1)
        helper.add_key_value_table(doc, [
            ("Actor chính", uc["actors"]),
            ("Actor hỗ trợ", uc["support"]),
            ("Kích hoạt", uc["trigger"]),
            ("Tiền điều kiện", " ".join(uc["pre"])),
            ("Hậu điều kiện", " ".join(uc["post"])),
        ])
        helper.add_section_heading(doc, "Luồng chính", 2)
        helper.add_numbered(doc, uc["main"])
        helper.add_section_heading(doc, "Luồng thay thế", 2)
        helper.add_bullets(doc, uc["alternative"])
        helper.add_section_heading(doc, "Luồng ngoại lệ", 2)
        helper.add_bullets(doc, uc["exception"])
        helper.add_section_heading(doc, "Quy tắc nghiệp vụ và validation", 2)
        helper.add_bullets(doc, uc["rules"])

    doc.add_page_break()
    helper.add_section_heading(doc, "14. Ghi chú nhất quán giữa các Use Case", 1)
    helper.add_bullets(doc, [
        "Mọi endpoint Admin phải xác minh danh tính, quyền thao tác và phạm vi dữ liệu ở backend.",
        "Các thao tác chuyển trạng thái, khóa tài khoản, duyệt đơn và xóa Customer phải có audit trail.",
        "Không xóa vật lý dữ liệu cần cho lịch sử đơn hàng, thanh toán, feedback, khảo sát hoặc báo cáo.",
        "Báo cáo chỉ dùng dữ liệu đã được xác nhận theo trạng thái nghiệp vụ chính thức và phải hiển thị phạm vi số liệu.",
        "Khi thay đổi trạng thái, schema hoặc quy tắc nghiệp vụ, phải cập nhật ERD/data dictionary và sequence tương ứng trước khi merge.",
    ])
    helper.add_paragraph(doc, "Tài liệu này là bộ đặc tả Admin tương ứng với các diagram trong thư mục docs/diagrams/admin và được thiết kế để dùng chung quy ước với đặc tả Customer.", italic=True)

    doc.core_properties.title = "Đặc tả Use Case Admin CRM và Sales"
    doc.core_properties.subject = "Part III Admin CRM and Sales"
    doc.core_properties.author = "Project Team"
    doc.core_properties.comments = "Admin use case specifications for MVP"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
